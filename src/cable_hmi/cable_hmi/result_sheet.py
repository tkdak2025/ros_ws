"""
검사 결과를 엑셀 파일(.xlsx)로 쓴다. '결과 파일 저장' 버튼 두 개(result_db.export_run / export_all)가 쓴다.

시트 세 개:
  검사 결과 - 사람이 보는 표. 한글 머리줄, 필요한 열만, 결과 칸 색(PASS 파랑 / FAIL 주황 / 그 밖 회색),
              첫 줄 고정, 필터, 숫자 소수 둘째 자리.
  요약      - 검사 1회(run) = 1행. 레시피, 시작 · 종료 시각, 끝난 이유, 집계.
  원본      - DB 의 열 전체(영어 필드 이름 그대로). 다시 분석할 때 쓴다.

openpyxl(python3-openpyxl) 이 없는 PC 에서는 예전처럼 CSV(엑셀용 BOM)로 대신 쓴다 - 저장 버튼이
라이브러리 하나 때문에 실패하지 않게. 어느 쪽으로 썼는지는 돌려주는 경로의 확장자로 알 수 있다.

Qt 에 의존하지 않는다.
"""

import csv
from pathlib import Path
from typing import Dict, List, Sequence

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:          # python3-openpyxl 이 없으면 CSV 로 대신 쓴다
    Workbook = None

# '검사 결과' 시트: (DB 열 이름, 머리줄, 종류). 종류 num = 소수 둘째 자리, time = 초까지 자른 시각.
RESULT_COLUMNS = (
    ('stamp', '검사 시간', 'time'),
    ('run_id', 'run', 'int'),
    ('recipe_id', '레시피', 'text'),
    ('point_id', 'Point', 'text'),
    ('point_name', '이름', 'text'),
    ('cable_type', '종류', 'text'),
    ('result', '결과', 'result'),
    ('reason_code', '사유 코드', 'text'),
    ('termination_reason', '종료 사유', 'text'),
    ('max_force_n', '최대 힘 (N)', 'num'),
    ('displacement_mm', '변위 (mm)', 'num'),
    ('soft_width_mm', 'Soft 기준 폭 (mm)', 'num'),
    ('pull_width_mm', 'Pull 최소 폭 (mm)', 'num'),
    ('width_delta_mm', '폭 변화 (mm)', 'num'),
    ('required_pull_force_n', '기준 힘 (N)', 'num'),
    ('displacement_limit_mm', '허용 변위 (mm)', 'num'),
    ('grip_failure_width_mm', '파지 실패 폭 (mm)', 'num'),
    ('reason', '판정 사유', 'text'),
)
RUN_COLUMNS = (
    ('run_id', 'run', 'int'),
    ('recipe_id', '레시피', 'text'),
    ('recipe_version', '버전', 'text'),
    ('started_at', '시작', 'time'),
    ('ended_at', '종료', 'time'),
    ('end_reason', '끝난 이유', 'text'),
    ('product_result', '제품 판정', 'text'),
    ('total_points', 'Point 수', 'int'),
    ('pass_count', 'PASS', 'int'),
    ('fail_count', 'FAIL', 'int'),
)

# 발표 자료 · HMI 와 같은 색 (PASS 파랑 / FAIL 주황)
_HEAD_FILL = 'E8F0FA'
_HEAD_FONT = '1B2530'
_RESULT_STYLE = {
    'PASS': ('1F5FAF', 'E8F0FA'),
    'FAIL': ('B5520F', 'FBEBDF'),
}
_OTHER_STYLE = ('5A6570', 'EEF1F4')
_LINE = 'DDE1E4'
# 열 너비 어림 (엑셀 너비 단위 = 숫자 '0' 한 글자). 굵은 글자 · 글꼴 차이를 덮으려고 넉넉히 잡는다.
_WIDTH_SCALE = 1.15
_CELL_PAD = 3            # 칸 양옆 여백
_FILTER_PAD = 3          # 머리줄 오른쪽의 필터 ▼ 단추가 글자를 가리지 않게
_MAX_WIDTH = 60          # 이보다 긴 글(판정 사유, 좌표 등)은 칸 안에서 줄을 바꾼다


def write(base: Path, rows: Sequence[Dict], runs: Sequence[Dict], raw_columns: List[str]) -> Path:
    """
    base(확장자 없는 경로) + .xlsx 로 쓰고 그 경로를 돌려준다. openpyxl 이 없으면 .csv.

    rows: 결과 1건 = dict(DB 열 이름 → 값). runs: 검사 1회 요약 dict 들. raw_columns: '원본' 시트 열 순서.
    쓰지 못하면 OSError 를 그대로 올린다(부르는 쪽이 사용자 메시지로 바꾼다).
    """
    base = Path(base)
    if Workbook is None:
        path = base.with_suffix('.csv')
        with path.open('w', encoding='utf-8-sig', newline='') as handle:
            writer = csv.writer(handle)
            writer.writerow(raw_columns)
            for row in rows:
                writer.writerow([row.get(c) for c in raw_columns])
        return path

    book = Workbook()
    sheet = book.active
    sheet.title = '검사 결과'
    _fill(sheet, RESULT_COLUMNS, rows)
    _fill(book.create_sheet('요약'), RUN_COLUMNS, runs)
    _fill(book.create_sheet('원본'), [(c, c, 'raw') for c in raw_columns], rows)
    path = base.with_suffix('.xlsx')
    book.save(path)
    return path


def _value(raw, kind):
    if raw is None or raw == '':
        return None
    if kind == 'time':
        text = str(raw).replace('T', ' ')
        return text[:19]                      # 2026-09-23 21:18:12 (시간대 표시는 뺀다)
    if kind == 'num':
        try:
            return float(raw)
        except (TypeError, ValueError):
            return raw
    if kind == 'int':
        try:
            return int(raw)
        except (TypeError, ValueError):
            return raw
    return raw


def _fill(sheet, columns, rows):
    side = Side(style='thin', color=_LINE)
    border = Border(left=side, right=side, top=side, bottom=side)
    widths = []
    for col, (_key, title, _kind) in enumerate(columns, start=1):
        cell = sheet.cell(row=1, column=col, value=title)
        cell.font = Font(bold=True, color=_HEAD_FONT)
        cell.fill = PatternFill('solid', fgColor=_HEAD_FILL)
        # 머리줄은 두 줄까지 접는다 - '파지 실패 폭 (mm)' 같은 긴 머리가 숫자 열을 넓히지 않게.
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = border
        widths.append(_header_width(title) + _FILTER_PAD)

    for r, row in enumerate(rows, start=2):
        for col, (key, _title, kind) in enumerate(columns, start=1):
            value = _value(row.get(key), kind)
            cell = sheet.cell(row=r, column=col, value=value)
            cell.border = border
            if kind == 'num' and isinstance(value, float):
                cell.number_format = '0.00'
                cell.alignment = Alignment(horizontal='right')
            elif kind in ('int', 'time', 'result'):
                cell.alignment = Alignment(horizontal='center')
                if kind == 'int' and isinstance(value, int):
                    cell.number_format = '0'     # run 번호(16자리)가 1.79E+15 로 보이지 않게
            if kind == 'result' and value:
                font, fill = _RESULT_STYLE.get(str(value), _OTHER_STYLE)
                cell.font = Font(bold=True, color=font)
                cell.fill = PatternFill('solid', fgColor=fill)
            if isinstance(value, float):
                shown = f'{value:.2f}'
            else:
                shown = '' if value is None else str(value)
            width = _text_width(shown)
            if width * _WIDTH_SCALE + _CELL_PAD > _MAX_WIDTH:
                # 너무 긴 글은 잘리지 않게 칸 안에서 줄을 바꾼다 (행 높이는 프로그램이 맞춘다).
                cell.alignment = Alignment(wrap_text=True, vertical='top')
            widths[col - 1] = max(widths[col - 1], width)

    for col, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(col)].width = min(
            _MAX_WIDTH, round(width * _WIDTH_SCALE) + _CELL_PAD)
    sheet.row_dimensions[1].height = 34          # 머리줄 두 줄
    sheet.freeze_panes = 'A2'
    # 인쇄 · PDF 로 뽑아도 열이 여러 쪽으로 잘리지 않게: 가로 방향, 너비는 한 쪽에 맞춘다.
    sheet.page_setup.orientation = 'landscape'
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    if rows:
        sheet.auto_filter.ref = sheet.dimensions


def _header_width(title: str) -> int:
    """머리를 빈칸에서 두 줄로 나눌 때 더 긴 줄의 너비가 가장 작아지는 값."""
    words = title.split(' ')
    best = _text_width(title)
    for i in range(1, len(words)):
        best = min(best, max(_text_width(' '.join(words[:i])), _text_width(' '.join(words[i:]))))
    return best


def _text_width(text: str) -> int:
    """엑셀 열 너비 어림값: 한글 등 넓은 글자는 2칸으로 센다."""
    return sum(2 if ord(ch) > 0x2E80 else 1 for ch in text)
